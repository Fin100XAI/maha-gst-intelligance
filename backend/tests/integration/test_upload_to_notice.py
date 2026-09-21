"""The whole road: a workbook in, a notice out.

This is the path the platform exists to walk, and every step of it was broken
in some way until it was walked end to end with a file that looks like one a
Commissionerate actually receives -- an accountant's export, not the
generator's output.

The five defects this test pins, each of which shipped looking correct:

* a credit note on a CDNR sheet was quarantined as a data error, because the
  document type was read from a hard-coded English header
* a credit note's amount was stored signed, so the engine negated it twice and
  *added* it to the demand
* Marathi tax columns did not map, so the tax became zero and the row still
  counted as parsed
* a GSTR-3B could not be ingested at all, and 21 of the 57 rules need one
* a period with no 3B on record was read as a 3B declaring nil, which reports
  the whole GSTR-1 liability as undeclared
"""

from __future__ import annotations

import io
from collections.abc import Iterator
from datetime import date
from decimal import Decimal

import pytest
from sqlalchemy import create_engine, select
from sqlalchemy.orm import Session, sessionmaker
from sqlalchemy.pool import StaticPool

from app.canonical import FinancialYear
from app.db.base import Base
from app.db.models import OutwardLine, Return3B, Taxpayer
from app.engine.load import gstins_in_snapshot, load_context
from app.engine.runner import run_for_taxpayer
from app.ingestion.persist import persist_report
from app.ingestion.pipeline import ingest_sheets
from app.ingestion.reader import read_workbook

FILER = "27AAGCS4521P1ZX"
BUYER = "27AABCU9603R1ZN"
OUT_OF_STATE = "29AACCB2894G1ZH"
AS_OF = date(2026, 9, 20)


@pytest.fixture
def db() -> Iterator[Session]:
    engine = create_engine(
        "sqlite://", future=True, connect_args={"check_same_thread": False}, poolclass=StaticPool
    )
    Base.metadata.create_all(engine)
    factory = sessionmaker(bind=engine, future=True, expire_on_commit=False)
    with factory() as session:
        yield session
    Base.metadata.drop_all(engine)
    engine.dispose()


def _gstr1() -> bytes:
    """An accountant's GSTR-1: title block, merged header, a credit note in
    parentheses, and a Marathi sheet."""
    from openpyxl import Workbook

    book = Workbook()
    b2b = book.active
    assert b2b is not None
    b2b.title = "B2B_072025"
    b2b.append(["GOODS AND SERVICES TAX - GSTR-1"])
    b2b.append([f"GSTIN {FILER}      Return Period: Jul-2025"])
    b2b.append([])
    b2b.append(
        [
            "GSTIN/UIN of Recipient",
            "Invoice Number",
            "Invoice date",
            "Place Of Supply",
            "Rate",
            "Taxable Value",
            "Integrated Tax Amount",
            "Central Tax Amount",
            "State/UT Tax Amount",
        ]
    )
    b2b.append(
        [BUYER, "INV/1", date(2025, 7, 4), "27-Maharashtra", 18, 1200000, None, 108000, 108000]
    )
    b2b.append(
        [OUT_OF_STATE, "INV/2", date(2025, 7, 9), "29-Karnataka", 18, 2200000, 396000, None, None]
    )

    cdnr = book.create_sheet("CDNR_072025")
    cdnr.append(["GOODS AND SERVICES TAX - GSTR-1"])
    cdnr.append([f"GSTIN {FILER}      Return Period: Jul-2025"])
    cdnr.append([])
    cdnr.append(
        [
            "GSTIN/UIN of Recipient",
            "Note Number",
            "Note date",
            "Note Type",
            "Place Of Supply",
            "Rate",
            "Taxable Value",
            "Central Tax Amount",
            "State/UT Tax Amount",
        ]
    )
    # Accountancy parentheses: a negative credit note.
    cdnr.append(
        [
            BUYER,
            "CN/1",
            date(2025, 7, 28),
            "C",
            "27-Maharashtra",
            18,
            "(2,00,000.00)",
            "(18,000.00)",
            "(18,000.00)",
        ]
    )

    marathi = book.create_sheet("B2B_082025_mr")
    marathi.append([f"जीएसटीआयएन {FILER}"])
    marathi.append([])
    marathi.append(
        [
            "प्राप्तकर्त्याचा जीएसटीआयएन",
            "बीजक क्रमांक",
            "बीजक दिनांक",
            "करपात्र मूल्य",
            "दर",
            "केंद्रीय कर",
            "राज्य कर",
        ]
    )
    marathi.append([BUYER, "INV/3", date(2025, 8, 6), 900000, 18, 81000, 81000])

    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def _gstr3b() -> bytes:
    """A GSTR-3B: the transposed summary table, under-declaring July."""
    from openpyxl import Workbook

    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.title = "GSTR3B_072025"
    sheet.append(["FORM GSTR-3B"])
    sheet.append([f"GSTIN {FILER}      Return Period: Jul-2025"])
    sheet.append([])
    sheet.append(
        [
            "Nature of Supplies",
            "Total Taxable value",
            "Integrated Tax",
            "Central Tax",
            "State/UT Tax",
            "Cess",
        ]
    )
    # GSTR-1 net of the credit note is IGST 396,000 / CGST 90,000 / SGST 90,000.
    # The 3B declares less on every head.
    sheet.append(
        [
            "(a) Outward taxable supplies (other than zero rated, nil rated and exempted)",
            3200000,
            196000,
            40000,
            40000,
            0,
        ]
    )
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def _ingest(session: Session, payload: bytes, name: str, snapshot_id: str | None) -> str:
    report = ingest_sheets(read_workbook(payload), filename=name, payload=payload)
    assert report.ledger.reconciles(), f"{name} did not reconcile"
    _, snapshot = persist_report(
        session,
        report,
        uploaded_by="sto.pune.1",
        storage_key=f"uploads/{name}",
        size_bytes=len(payload),
        snapshot_id=snapshot_id,
    )
    session.flush()
    return snapshot.id


class TestUploadToFinding:
    @pytest.mark.golden
    def test_the_whole_road(self, db: Session) -> None:
        snapshot = _ingest(db, _gstr1(), "gstr1.xlsx", None)
        _ingest(db, _gstr3b(), "gstr3b.xlsx", snapshot)

        # --- ingestion registered the filer, so it is not invisible ---------
        registered = db.get(Taxpayer, FILER)
        assert registered is not None
        assert registered.status == "FROM_RETURN"
        assert gstins_in_snapshot(db, snapshot) == [FILER]

        # --- the credit note survived, as a magnitude ----------------------
        note = (
            db.execute(select(OutwardLine).where(OutwardLine.doc_type == "CREDIT_NOTE"))
            .scalars()
            .one()
        )
        assert note.taxable_value == Decimal("200000.00")
        assert note.cgst == Decimal("18000.00"), "a credit note is stored as a magnitude"

        # --- the Marathi tax columns mapped --------------------------------
        august = (
            db.execute(select(OutwardLine).where(OutwardLine.period == "082025")).scalars().one()
        )
        assert august.cgst == Decimal("81000.00")
        assert august.sgst == Decimal("81000.00")

        # --- the 3B was ingested at all ------------------------------------
        three_b = db.execute(select(Return3B)).scalars().one()
        assert three_b.period == "072025"
        assert three_b.t31a_igst == Decimal("196000.00")

        # --- the engine ran over what was ingested -------------------------
        context = load_context(db, FILER, snapshot_id=snapshot, fy=FinancialYear(2025), as_of=AS_OF)
        outcome = run_for_taxpayer(context, with_identities=True)

        triggered = [f for f in outcome.findings if f.triggered]
        assert [f.rule_id for f in triggered] == ["OUT-01"]

        # GSTR-1 net of the credit note, less what the 3B declared:
        #   IGST 396,000 - 196,000 = 200,000
        #   CGST  90,000 -  40,000 =  50,000
        finding = triggered[0]
        assert finding.delta.igst == Decimal("200000.00")
        assert finding.delta.cgst == Decimal("50000.00")
        assert finding.delta.sgst == Decimal("50000.00")

    @pytest.mark.golden
    def test_a_period_without_a_3b_is_not_read_as_a_nil_return(self, db: Session) -> None:
        """August has a GSTR-1 and no 3B. Reading the absence as a nil
        declaration would report the whole August liability as undeclared and
        route a notice for a discrepancy nobody has established."""
        snapshot = _ingest(db, _gstr1(), "gstr1.xlsx", None)
        _ingest(db, _gstr3b(), "gstr3b.xlsx", snapshot)

        context = load_context(db, FILER, snapshot_id=snapshot, fy=FinancialYear(2025), as_of=AS_OF)
        outcome = run_for_taxpayer(context, with_identities=True)

        august = [
            row
            for rows in outcome.identities.values()
            for row in rows
            if row.identity_id == "R1" and row.period is not None and row.period.mmyyyy == "082025"
        ]
        assert august, "R1 was not evaluated for August at all"
        assert august[0].status == "NOT_EVALUATED"
        assert "gstr3b" in " ".join(august[0].missing_inputs)

        # And no finding claims an August discrepancy.
        assert not [
            f
            for f in outcome.findings
            if f.triggered and f.period is not None and f.period.mmyyyy == "082025"
        ]

    def test_the_matrix_names_the_return_each_period_lacks(self, db: Session) -> None:
        snapshot = _ingest(db, _gstr1(), "gstr1.xlsx", None)
        _ingest(db, _gstr3b(), "gstr3b.xlsx", snapshot)

        context = load_context(db, FILER, snapshot_id=snapshot, fy=FinancialYear(2025), as_of=AS_OF)
        outcome = run_for_taxpayer(context, with_identities=True)

        r1 = [
            row for rows in outcome.identities.values() for row in rows if row.identity_id == "R1"
        ]
        breached = [row for row in r1 if row.status == "BREACHED"]
        unevaluated = [row for row in r1 if row.status == "NOT_EVALUATED"]

        assert len(breached) == 1, "only July can be tested"
        assert breached[0].period is not None
        assert breached[0].period.mmyyyy == "072025"
        assert unevaluated, "every other period should name the return it lacks"
        for row in unevaluated:
            assert row.missing_inputs, f"{row.period} gives no reason"


class TestNoticeGate:
    @pytest.mark.golden
    def test_a_stub_registration_cannot_be_served(self, db: Session) -> None:
        """A notice addressed to a GSTIN rather than to a person is defective.

        Ingestion knows the filer only from the return, which carries no legal
        name. That is enough to analyse a taxpayer and not enough to serve one.
        """
        snapshot = _ingest(db, _gstr1(), "gstr1.xlsx", None)
        registered = db.get(Taxpayer, FILER)
        assert registered is not None
        assert registered.legal_name == registered.gstin
        assert registered.status == "FROM_RETURN"
        assert snapshot
