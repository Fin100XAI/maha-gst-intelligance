"""Phase 1 acceptance: the pipeline survives real files and never loses a row."""

from __future__ import annotations

import io
from datetime import date
from decimal import Decimal
from typing import Any

import pytest
from openpyxl import Workbook

from app.ingestion.coerce import CoercionError, coerce_date, coerce_hsn, coerce_rate
from app.ingestion.header import detect_header, merge_header_rows
from app.ingestion.pipeline import (
    duplicate_key,
    ingest_sheet,
    ingest_sheets,
    snapshot_hash,
)
from app.ingestion.quarantine import QuarantineReason
from app.ingestion.reader import RawSheet, read_workbook
from app.ingestion.sniffer import classify_sheet
from app.ingestion.synonyms import match_header, match_headers

OWNER = "27AAPFU0939F1ZV"
SUPPLIER = "27AACCM9910C1ZN"  # checksum-valid, synthetic

# The GST portal's own GSTR-1 B2B export header, verbatim.
PORTAL_B2B = [
    "GSTIN/UIN of Recipient",
    "Receiver Name",
    "Invoice Number",
    "Invoice date",
    "Invoice Value",
    "Place Of Supply",
    "Reverse Charge",
    "Applicable % of Tax Rate",
    "Invoice Type",
    "E-Commerce GSTIN",
    "Rate",
    "Taxable Value",
    "Cess Amount",
]


def b2b_row(doc_no: str, taxable: str, rate: str = "18") -> list[Any]:
    return [
        SUPPLIER,
        "Meridian Traders",
        doc_no,
        "15-06-2025",
        None,
        "27-Maharashtra",
        "N",
        None,
        "Regular B2B",
        None,
        rate,
        taxable,
        "0",
    ]


def sheet_from(rows: list[list[Any]], name: str = "b2b", index: int = 0) -> RawSheet:
    return RawSheet(name=name, index=index, rows=rows)


# ---------------------------------------------------------------------------
# synonyms
# ---------------------------------------------------------------------------


class TestSynonyms:
    def test_portal_headers_map(self) -> None:
        matches = {m.header: m.field for m in match_headers(list(PORTAL_B2B), "GSTR1")}
        assert matches["GSTIN/UIN of Recipient"] == "counterparty_gstin"
        assert matches["Invoice Number"] == "doc_no"
        assert matches["Invoice date"] == "doc_date"
        assert matches["Place Of Supply"] == "pos"
        assert matches["Taxable Value"] == "taxable_value"
        assert matches["Rate"] == "rate"

    @pytest.mark.parametrize(
        ("header", "field"),
        [
            ("Integrated Tax(₹)", "igst"),
            ("Central Tax (Rs.)", "cgst"),
            ("State/UT Tax", "sgst"),
            ("Cess(₹)", "cess"),
            ("Taxable Value (₹)", "taxable_value"),
            ("HSN/SAC", "hsn"),
        ],
    )
    def test_currency_ornament_is_stripped(self, header: str, field: str) -> None:
        assert match_header(header, "GSTR1").field == field

    @pytest.mark.golden
    @pytest.mark.parametrize(
        ("header", "field"),
        [
            ("पावती क्रमांक", "doc_no"),
            ("पावती दिनांक", "doc_date"),
            ("करपात्र मूल्य", "taxable_value"),
            ("दर", "rate"),
        ],
    )
    def test_marathi_headers_map_from_day_one(self, header: str, field: str) -> None:
        assert match_header(header, "GSTR1").field == field

    def test_misspellings_fall_through_to_edit_distance(self) -> None:
        match = match_header("Taxabel Value", "GSTR1")
        assert match.field == "taxable_value"
        assert match.method == "fuzzy"
        assert match.needs_confirmation is True

    def test_a_field_is_never_claimed_twice(self) -> None:
        matches = match_headers(["Taxable Value", "Taxable Value (subtotal)"], "GSTR1")
        assigned = [m.field for m in matches if m.field == "taxable_value"]
        assert len(assigned) == 1
        assert any(m.method == "duplicate_field" for m in matches)

    def test_nonsense_is_reported_not_guessed(self) -> None:
        assert match_header("Remarks by the auditor", "GSTR1").field is None


# ---------------------------------------------------------------------------
# coercion
# ---------------------------------------------------------------------------


class TestCoercion:
    @pytest.mark.golden
    @pytest.mark.parametrize(
        ("raw", "expected"),
        [
            ("15-06-2025", date(2025, 6, 15)),
            ("15/06/2025", date(2025, 6, 15)),
            ("15/06/25", date(2025, 6, 15)),
            ("2025-06-15", date(2025, 6, 15)),
            ("15-Jun-2025", date(2025, 6, 15)),
            ("15 June 2025", date(2025, 6, 15)),
            (date(2025, 6, 15), date(2025, 6, 15)),
            (Decimal("45823"), date(2025, 6, 15)),  # Excel serial
        ],
    )
    def test_every_date_spelling(self, raw: object, expected: date) -> None:
        assert coerce_date(raw) == expected

    def test_ambiguous_numeric_dates_are_read_day_first(self) -> None:
        """The portal and every Indian accounting package write dd-mm-yyyy."""
        assert coerce_date("06-07-2025") == date(2025, 7, 6)

    def test_an_unreadable_date_raises_rather_than_becoming_none(self) -> None:
        with pytest.raises(CoercionError):
            coerce_date("the fifteenth")

    @pytest.mark.parametrize(
        ("raw", "expected"),
        [("18", "18"), ("18%", "18"), ("0.18", "18"), (Decimal("5"), "5"), ("0.25", "0.25")],
    )
    def test_rate_conventions(self, raw: object, expected: str) -> None:
        """Some exports write 0.18 for 18%.  The rough-diamond 0.25% rate is
        the one value that means itself in both conventions."""
        assert coerce_rate(raw) == Decimal(expected)

    def test_a_rate_above_100_is_refused(self) -> None:
        with pytest.raises(CoercionError):
            coerce_rate("180")

    @pytest.mark.parametrize(
        ("raw", "expected"),
        [
            ("0902", "0902"),
            (Decimal("902"), "0902"),
            ("998314", "998314"),
            ("85", "85"),
            ("8517", "8517"),
        ],
    )
    def test_hsn_is_padded_to_a_valid_length(self, raw: object, expected: str) -> None:
        """Excel turns 0902 into the number 902; the leading zero must come back."""
        assert coerce_hsn(raw) == expected


# ---------------------------------------------------------------------------
# header detection
# ---------------------------------------------------------------------------


class TestHeaderDetection:
    @pytest.mark.golden
    def test_finds_the_header_behind_a_four_row_title_block(self) -> None:
        rows: list[list[Any]] = [
            ["COMMERCIAL TAXES DEPARTMENT", None, None, None],
            ["GSTR-1 Outward Supplies", None, None, None],
            ["FY 2025-26 · June", None, None, None],
            [None, None, None, None],
            PORTAL_B2B,
            *[b2b_row(f"INV-{i}", "100000") for i in range(3)],
        ]
        detection = detect_header(rows, "GSTR1")
        assert detection is not None
        assert detection.row_index == 4  # the fifth row, behind a four-row block
        assert detection.data_starts_at == 5

    def test_merges_a_two_row_header(self) -> None:
        labels = merge_header_rows(
            [
                [None, "Invoice", None, "Tax", None, None],
                ["GSTIN", "Number", "Date", "IGST", "CGST", "SGST"],
            ]
        )
        assert labels == [
            "GSTIN",
            "Invoice Number",
            "Invoice Date",
            "Tax IGST",
            "Tax CGST",
            "Tax SGST",
        ]

    def test_a_sheet_with_no_header_returns_none(self) -> None:
        assert detect_header([], "GSTR1") is None


# ---------------------------------------------------------------------------
# classification
# ---------------------------------------------------------------------------


class TestClassification:
    @pytest.mark.parametrize(
        ("name", "headers", "family"),
        [
            ("b2b", PORTAL_B2B, "GSTR1"),
            ("B2CS(7)", PORTAL_B2B, "GSTR1"),
            (
                "GSTR2B_Jun2025",
                ["GSTIN of supplier", "Invoice number", "Taxable Value", "ITC Availability"],
                "GSTR2B",
            ),
            (
                "EWB Report",
                ["E Way Bill No", "Vehicle No", "Distance", "Total Inv Value"],
                "EWAYBILL",
            ),
            ("IRN list", ["IRN", "Ack No", "Ack Date", "Taxable Value"], "EINVOICE"),
            (
                "Credit Ledger",
                ["Date", "Opening Balance", "Credit", "Debit", "Closing Balance"],
                "LEDGER",
            ),
        ],
    )
    def test_families(self, name: str, headers: list[str], family: str) -> None:
        assert classify_sheet(name, list(headers), []).family == family

    def test_sections_come_from_the_sheet_name(self) -> None:
        assert classify_sheet("b2b", PORTAL_B2B, []).section == "B2B"
        assert classify_sheet("cdnr", PORTAL_B2B, []).section == "CDNR"
        assert classify_sheet("b2ba", PORTAL_B2B, []).section == "AMENDMENT"

    def test_an_unrecognised_sheet_says_so_rather_than_guessing(self) -> None:
        result = classify_sheet("Notes", ["Remarks", "Prepared by", "Signature"], [])
        assert result.family is None
        assert result.needs_confirmation is True

    def test_value_shapes_classify_a_marathi_headed_sheet(self) -> None:
        """A column of GSTIN-shaped strings is evidence even when the header
        is in a script the lexicon has not seen."""
        marathi = [
            "पुरवठादार जीएसटीआयएन",
            "करपात्र मूल्य",
        ]
        rows = [[SUPPLIER, Decimal("100000")]]
        assert classify_sheet("purchase register", marathi, rows).family == "GSTR2B"


# ---------------------------------------------------------------------------
# the row identity -- Law 5
# ---------------------------------------------------------------------------


class TestRowIdentity:
    @pytest.mark.golden
    def test_rows_in_equals_parsed_plus_quarantined_plus_duplicates(self) -> None:
        rows: list[list[Any]] = [
            PORTAL_B2B,
            b2b_row("INV-1", "100000"),
            b2b_row("INV-2", "200000"),
            b2b_row("INV-2", "200000"),  # an exact duplicate
            [*b2b_row("INV-3", "100000")[:3], "not a date", *b2b_row("INV-3", "100000")[4:]],
            ["Total", None, None, None, None, None, None, None, None, None, None, "300000", None],
        ]
        outcome = ingest_sheet(sheet_from(rows), owner_gstin=OWNER, period_hint="062025")
        ledger = outcome.ledger
        assert ledger.rows_in == 5
        assert ledger.parsed == 2
        assert ledger.duplicates == 1
        assert ledger.quarantined == 2  # the bad date and the totals row
        assert ledger.reconciles() is True

    def test_a_quarantined_row_carries_its_reason_and_its_original_cells(self) -> None:
        bad = b2b_row("INV-9", "100000")
        bad[3] = "not a date"
        outcome = ingest_sheet(
            sheet_from([PORTAL_B2B, bad]), owner_gstin=OWNER, period_hint="062025"
        )
        assert outcome.ledger.quarantined == 1
        row = outcome.ledger.quarantined_rows[0]
        assert row.reason_code is QuarantineReason.COERCION_FAILED
        assert "date" in row.reason
        assert row.original_cells["Invoice date"] == "not a date"
        assert row.original_cells["Invoice Number"] == "INV-9"

    @pytest.mark.golden
    def test_a_line_whose_tax_contradicts_its_rate_is_quarantined(self) -> None:
        """taxable x rate must equal the tax charged within a rupee."""
        headers = [*PORTAL_B2B, "Integrated Tax"]
        good = [*b2b_row("INV-1", "100000"), "18000"]
        bad = [*b2b_row("INV-2", "100000"), "12000"]  # 18% of 1,00,000 is 18,000
        outcome = ingest_sheet(
            sheet_from([headers, good, bad]), owner_gstin=OWNER, period_hint="062025"
        )
        assert outcome.ledger.parsed == 1
        assert outcome.ledger.quarantined == 1
        quarantined = outcome.ledger.quarantined_rows[0]
        assert quarantined.reason_code is QuarantineReason.CROSS_FIELD_MISMATCH
        assert "18000.00" in quarantined.reason

    def test_a_failed_gstin_checksum_is_quarantined_with_that_reason(self) -> None:
        bad = b2b_row("INV-1", "100000")
        bad[0] = "27AACCM9910C1ZM"  # structurally valid, wrong checksum
        outcome = ingest_sheet(
            sheet_from([PORTAL_B2B, bad]), owner_gstin=OWNER, period_hint="062025"
        )
        assert outcome.ledger.quarantined == 1
        assert outcome.ledger.quarantined_rows[0].reason_code is QuarantineReason.GSTIN_INVALID
        assert "checksum" in outcome.ledger.quarantined_rows[0].reason

    def test_blank_spacer_rows_are_not_counted_as_rows(self) -> None:
        rows: list[list[Any]] = [
            PORTAL_B2B,
            b2b_row("INV-1", "100000"),
            [None] * len(PORTAL_B2B),
            b2b_row("INV-2", "100000"),
        ]
        outcome = ingest_sheet(sheet_from(rows), owner_gstin=OWNER, period_hint="062025")
        assert outcome.ledger.rows_in == 2
        assert outcome.ledger.parsed == 2


# ---------------------------------------------------------------------------
# duplicates and snapshots
# ---------------------------------------------------------------------------


class TestIdempotence:
    def _workbook_bytes(self) -> bytes:
        workbook = Workbook()
        sheet = workbook.active
        assert sheet is not None
        sheet.title = "b2b"
        sheet.append(PORTAL_B2B)
        for index in range(5):
            sheet.append(b2b_row(f"INV-{index}", "100000"))
        buffer = io.BytesIO()
        workbook.save(buffer)
        return buffer.getvalue()

    @pytest.mark.golden
    def test_re_uploading_the_same_file_is_a_non_event(self) -> None:
        payload = self._workbook_bytes()

        first = ingest_sheets(
            read_workbook(payload),
            filename="GSTR1_Jun2025.xlsx",
            payload=payload,
            owner_gstin=OWNER,
            period_hint="062025",
        )
        assert first.ledger.parsed == 5
        assert first.ledger.duplicates == 0

        known = {duplicate_key(record) for record in first.records}
        second = ingest_sheets(
            read_workbook(payload),
            filename="GSTR1_Jun2025.xlsx",
            payload=payload,
            owner_gstin=OWNER,
            period_hint="062025",
            known_keys=known,
        )
        assert second.ledger.parsed == 0
        assert second.ledger.duplicates == 5
        assert second.ledger.reconciles() is True

    def test_the_snapshot_hash_does_not_depend_on_row_order(self) -> None:
        payload = self._workbook_bytes()
        report = ingest_sheets(
            read_workbook(payload),
            filename="a.xlsx",
            payload=payload,
            owner_gstin=OWNER,
            period_hint="062025",
        )
        forwards = snapshot_hash(report.records)
        backwards = snapshot_hash(list(reversed(report.records)))
        assert forwards == backwards

    def test_the_duplicate_key_is_content_not_position(self) -> None:
        one = ingest_sheet(
            sheet_from([PORTAL_B2B, b2b_row("INV-1", "100000")]),
            owner_gstin=OWNER,
            period_hint="062025",
        ).records[0]
        other = ingest_sheet(
            sheet_from([[None] * 13, PORTAL_B2B, b2b_row("INV-1", "100000")], name="elsewhere"),
            owner_gstin=OWNER,
            period_hint="062025",
        ).records[0]
        assert duplicate_key(one) == duplicate_key(other)


# ---------------------------------------------------------------------------
# the Phase 1 gate: a 40-tab workbook
# ---------------------------------------------------------------------------


def build_messy_workbook() -> bytes:
    """A workbook shaped like the ones officers actually upload."""
    workbook = Workbook()
    first = workbook.active
    assert first is not None
    first.title = "b2b"
    # A logo row, a title block, and then the real header on row 5.
    first.append(["COMMERCIAL TAXES DEPARTMENT"])
    first.append(["GSTR-1 Outward Supplies"])
    first.append(["FY 2025-26 · June"])
    first.append([])
    first.append(PORTAL_B2B)
    for index in range(10):
        first.append(b2b_row(f"INV-{index:03d}", "100000"))

    # 38 further recognisable sheets.
    for index in range(38):
        sheet = workbook.create_sheet(f"b2b_part{index}")
        sheet.append(PORTAL_B2B)
        sheet.append(b2b_row(f"P{index}-1", "50000"))

    # One sheet nobody can read.
    unreadable = workbook.create_sheet("Notes")
    unreadable.append(["Prepared by", "Checked by", "Remarks"])
    unreadable.append(["R. Deshmukh", "S. Kulkarni", "see covering letter"])

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


class TestPhase1Gate:
    @pytest.mark.golden
    def test_forty_tabs_one_unreadable_and_the_other_thirty_nine_ingest(self) -> None:
        payload = build_messy_workbook()
        sheets = read_workbook(payload)
        assert len(sheets) == 40

        report = ingest_sheets(
            sheets,
            filename="messy.xlsx",
            payload=payload,
            owner_gstin=OWNER,
            period_hint="062025",
        )

        parsed_sheets = [s for s in report.sheets if s.status == "PARSED"]
        unrecognised = [s for s in report.sheets if s.status == "UNRECOGNISED"]

        assert len(parsed_sheets) == 39, "one bad sheet must not reject the workbook"
        assert len(unrecognised) == 1
        assert unrecognised[0].sheet_name == "Notes"
        assert (
            unrecognised[0].ledger.quarantined_rows[0].reason_code
            is QuarantineReason.SHEET_UNRECOGNISED
        )

        assert report.ledger.parsed == 10 + 38
        assert report.ledger.reconciles() is True

    def test_the_header_is_found_behind_the_title_block(self) -> None:
        payload = build_messy_workbook()
        report = ingest_sheets(
            read_workbook(payload),
            filename="messy.xlsx",
            payload=payload,
            owner_gstin=OWNER,
            period_hint="062025",
        )
        b2b = next(sheet for sheet in report.sheets if sheet.sheet_name == "b2b")
        assert b2b.header_row == 4
        assert b2b.ledger.parsed == 10

    def test_the_report_carries_everything_the_screen_needs(self) -> None:
        payload = build_messy_workbook()
        report = ingest_sheets(
            read_workbook(payload),
            filename="messy.xlsx",
            payload=payload,
            owner_gstin=OWNER,
            period_hint="062025",
        )
        body = report.as_dict()
        assert body["counts"]["reconciles"] is True
        assert body["file_sha256"]
        assert len(body["sheets"]) == 40
        assert body["quarantine"][0]["reason"]
